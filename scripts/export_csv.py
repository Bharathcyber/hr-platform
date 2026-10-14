
import csv
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = BASE_DIR / "01_LedgerForge_Python_Master_Dataset.xlsx"
EXPORT_DIR = BASE_DIR / "exports"

SHEETS_TO_EXPORT = [
    "employees",
    "salary_structures",
    "attendance",
]


def export_sheet(workbook, sheet_name):
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)

    headers = next(rows)
    headers = [
        str(value).strip() if value is not None else ""
        for value in headers
    ]

    output_path = EXPORT_DIR / f"{sheet_name}.csv"
    data_count = 0

    with output_path.open(
        "w", newline="", encoding="utf-8-sig"
    ) as file:
        writer = csv.writer(file)
        writer.writerow(headers)

        for row in rows:
            if not any(value is not None for value in row):
                continue

            writer.writerow([
                value.isoformat()
                if hasattr(value, "isoformat")
                else value
                for value in row
            ])
            data_count += 1

    print(f"{sheet_name}: {data_count} rows -> {output_path}")
    return data_count


def main():
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(
        WORKBOOK_PATH,
        read_only=True,
        data_only=True,
    )
    try:
        for sheet_name in SHEETS_TO_EXPORT:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Required workbook sheet missing: {sheet_name}"
                )
            export_sheet(workbook, sheet_name)
    finally:
        workbook.close()

    print("CSV export completed successfully.")


if __name__ == "__main__":
    main()
