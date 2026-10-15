import argparse
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.dialects.postgresql import insert

from app.core.config import settings
from app.models import (
    Department,
    Employee,
    EmployeeAssignment,
    Holiday,
    LeaveType,
    Role,
    SalaryComponent,
    User,
)


BASE_DIR = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = (
    BASE_DIR / "data" / "seed" / "01_LedgerForge_Python_Master_Dataset.xlsx"
)


def read_sheet(workbook, sheet_name):
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)

    try:
        first_row = next(rows)
    except StopIteration:
        raise ValueError(f"{sheet_name}: worksheet is empty")

    headers = [
        str(value).strip() if value is not None else ""
        for value in first_row
    ]

    if not all(headers) or len(headers) != len(set(headers)):
        raise ValueError(f"{sheet_name}: invalid or duplicate headers")

    records = []
    for values in rows:
        if not any(value is not None for value in values):
            continue
        records.append(dict(zip(headers, values)))

    return records


def required_text(value, field, sheet):
    if value is None or not str(value).strip():
        raise ValueError(f"{sheet}: missing required value for {field}")
    return str(value).strip()


def optional_text(value):
    if value is None or not str(value).strip():
        return None
    return str(value).strip()


def as_bool(value, field, sheet):
    if isinstance(value, bool):
        return value

    if value is None:
        raise ValueError(f"{sheet}: missing required boolean {field}")

    normalized = str(value).strip().upper()

    if normalized in {"Y", "YES", "TRUE", "1"}:
        return True
    if normalized in {"N", "NO", "FALSE", "0"}:
        return False

    raise ValueError(
        f"{sheet}: invalid boolean value for {field}: {value!r}"
    )


def as_date(value, field, sheet, required=False):
    if value is None:
        if required:
            raise ValueError(f"{sheet}: missing required date {field}")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"{sheet}: invalid date in {field}: {value!r}")


def as_datetime(value, field, sheet, required=False):
    if value is None:
        if required:
            raise ValueError(f"{sheet}: missing required datetime {field}")
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    raise ValueError(f"{sheet}: invalid datetime in {field}: {value!r}")


def as_decimal(value, field, sheet):
    if value is None:
        return None
    try:
        return Decimal(str(value))
    except Exception as exc:
        raise ValueError(
            f"{sheet}: invalid numeric value for {field}: {value!r}"
        ) from exc


def as_int(value, field, sheet, required=False, default=None):
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ValueError(f"{sheet}: missing required integer {field}")
        return default
    try:
        number = Decimal(str(value))
        if number != number.to_integral_value():
            raise ValueError
        return int(number)
    except Exception as exc:
        raise ValueError(
            f"{sheet}: invalid integer value for {field}: {value!r}"
        ) from exc


def unique_ids(rows, field, sheet):
    seen = set()
    for row in rows:
        value = required_text(row.get(field), field, sheet)
        if value in seen:
            raise ValueError(f"{sheet}: duplicate {field} {value!r}")
        seen.add(value)
    return seen


def load_and_validate():
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    workbook = load_workbook(
        WORKBOOK_PATH, read_only=True, data_only=True
    )

    try:
        departments_raw = read_sheet(workbook, "departments")
        roles_raw = read_sheet(workbook, "roles")
        employees_raw = read_sheet(workbook, "employees")
        assignments_raw = read_sheet(workbook, "employee_assignments")
        users_raw = read_sheet(workbook, "users")
        salary_components_raw = read_sheet(workbook, "salary_components")
        holidays_raw = read_sheet(workbook, "holidays")
        leave_types_raw = read_sheet(workbook, "leave_types")
    finally:
        workbook.close()

    department_ids = unique_ids(
        departments_raw, "department_id", "departments"
    )
    role_ids = unique_ids(roles_raw, "role_id", "roles")
    employee_ids = unique_ids(employees_raw, "employee_id", "employees")
    unique_ids(assignments_raw, "assignment_id", "employee_assignments")
    unique_ids(users_raw, "user_id", "users")
    unique_ids(
        salary_components_raw, "component_id", "salary_components"
    )
    unique_ids(leave_types_raw, "leave_type_code", "leave_types")

    holiday_dates = [
        as_date(
            row.get("holiday_date"),
            "holiday_date",
            "holidays",
            required=True,
        )
        for row in holidays_raw
    ]
    if len(holiday_dates) != len(set(holiday_dates)):
        raise ValueError("holidays: duplicate holiday_date values")

    departments = []
    for row in departments_raw:
        departments.append({
            "department_id": required_text(
                row.get("department_id"), "department_id", "departments"
            ),
            "department_name": required_text(
                row.get("department_name"), "department_name", "departments"
            ),
            "cost_center": optional_text(row.get("cost_center")),
            "description": optional_text(row.get("description")),
        })

    roles = []
    for row in roles_raw:
        role_id = required_text(row.get("role_id"), "role_id", "roles")
        department_id = required_text(
            row.get("department_id"), "department_id", "roles"
        )

        if department_id not in department_ids:
            raise ValueError(
                f"Role {role_id} references unknown department {department_id}"
            )

        overtime = str(row.get("overtime_eligible") or "").strip().upper()
        if overtime not in {"Y", "N", "TRUE", "FALSE"}:
            raise ValueError(
                f"Role {role_id}: invalid overtime_eligible {overtime!r}"
            )

        roles.append({
            "role_id": role_id,
            "department_id": department_id,
            "role_title": required_text(
                row.get("role_title"), "role_title", "roles"
            ),
            "level": optional_text(row.get("level")),
            "grade": optional_text(row.get("grade")),
            "band_min_annual": as_decimal(
                row.get("band_min_annual"), "band_min_annual", "roles"
            ),
            "band_mid_annual": as_decimal(
                row.get("band_mid_annual"), "band_mid_annual", "roles"
            ),
            "band_max_annual": as_decimal(
                row.get("band_max_annual"), "band_max_annual", "roles"
            ),
            "overtime_eligible": overtime in {"Y", "TRUE"},
            "band_year": as_int(
                row.get("band_year"), "band_year", "roles"
            ),
        })

    employees = []
    for row in employees_raw:
        employee_id = required_text(
            row.get("employee_id"), "employee_id", "employees"
        )
        department_id = optional_text(row.get("department_id"))
        role_id = optional_text(row.get("role_id"))
        manager_id = optional_text(row.get("manager_id"))

        if department_id and department_id not in department_ids:
            raise ValueError(
                f"Employee {employee_id}: unknown department {department_id}"
            )
        if role_id and role_id not in role_ids:
            raise ValueError(
                f"Employee {employee_id}: unknown role {role_id}"
            )
        if manager_id and manager_id not in employee_ids:
            raise ValueError(
                f"Employee {employee_id}: unknown manager {manager_id}"
            )

        hire_date = as_date(
            row.get("hire_date"), "hire_date", "employees", required=True
        )
        exit_date = as_date(row.get("exit_date"), "exit_date", "employees")
        if exit_date and exit_date < hire_date:
            raise ValueError(
                f"Employee {employee_id}: exit_date is before hire_date"
            )

        pf = optional_text(row.get("pf_applicable"))
        esi = optional_text(row.get("esi_applicable"))

        if pf is not None:
            pf = pf.upper()
            if pf not in {"Y", "N"}:
                raise ValueError(f"Employee {employee_id}: invalid PF value")
        if esi is not None:
            esi = esi.upper()
            if esi not in {"Y", "N"}:
                raise ValueError(f"Employee {employee_id}: invalid ESI value")


        experience = as_decimal(
            row.get("prior_experience_years"),
            "prior_experience_years",
            "employees",
    )
        if experience is None:
            experience = Decimal("0")

        career_gap = as_int(
            row.get("career_gap_months"),
            "career_gap_months",
            "employees",
            default=0,
        )
        notice_period = as_int(
            row.get("notice_period_days"),
            "notice_period_days",
            "employees",
        )

        if experience < 0 or career_gap < 0:
            raise ValueError(
                f"Employee {employee_id}: negative experience or career gap"
            )
        if notice_period is not None and notice_period < 0:
            raise ValueError(
                f"Employee {employee_id}: negative notice period"
            )

        employees.append({
            "employee_id": employee_id,
            "first_name": required_text(
                row.get("first_name"), "first_name", "employees"
            ),
            "last_name": required_text(
                row.get("last_name"), "last_name", "employees"
            ),
            "full_name": required_text(
                row.get("full_name"), "full_name", "employees"
            ),
            "work_email": required_text(
                row.get("work_email"), "work_email", "employees"
            ),
            "gender": optional_text(row.get("gender")),
            "date_of_birth": as_date(
                row.get("date_of_birth"), "date_of_birth", "employees"
            ),
            "hire_date": hire_date,
            "exit_date": exit_date,
            "employment_status": required_text(
                row.get("employment_status"), "employment_status", "employees"
            ),
            "employment_type": optional_text(row.get("employment_type")),
            "department_id": department_id,
            "role_id": role_id,
            "manager_id": manager_id,
            "work_location": optional_text(row.get("work_location")),
            "work_mode": optional_text(row.get("work_mode")),
            "highest_education": optional_text(row.get("highest_education")),
            "prior_experience_years": experience,
            "career_gap_months": career_gap,
            "commute_distance_km": as_decimal(
                row.get("commute_distance_km"),
                "commute_distance_km",
                "employees",
            ),
            "pf_applicable": pf,
            "esi_applicable": esi,
            "notice_period_days": notice_period,
            "resignation_date": as_date(
                row.get("resignation_date"), "resignation_date", "employees"
            ),
            "exit_type": optional_text(row.get("exit_type")),
            "exit_reason": optional_text(row.get("exit_reason")),
            "created_at": as_datetime(
                row.get("created_at"), "created_at", "employees"
            ),
        })

    assignments = []
    for row in assignments_raw:
        assignment_id = required_text(
            row.get("assignment_id"), "assignment_id", "employee_assignments"
        )
        employee_id = required_text(
            row.get("employee_id"), "employee_id", "employee_assignments"
        )
        department_id = optional_text(row.get("department_id"))
        role_id = optional_text(row.get("role_id"))
        manager_id = optional_text(row.get("manager_id"))

        if employee_id not in employee_ids:
            raise ValueError(
                f"Assignment {assignment_id}: unknown employee {employee_id}"
            )
        if department_id and department_id not in department_ids:
            raise ValueError(
                f"Assignment {assignment_id}: unknown department {department_id}"
            )
        if role_id and role_id not in role_ids:
            raise ValueError(
                f"Assignment {assignment_id}: unknown role {role_id}"
            )
        if manager_id and manager_id not in employee_ids:
            raise ValueError(
                f"Assignment {assignment_id}: unknown manager {manager_id}"
            )

        start = as_datetime(
            row.get("effective_from"),
            "effective_from",
            "employee_assignments",
            required=True,
        )
        end = as_datetime(
            row.get("effective_to"),
            "effective_to",
            "employee_assignments",
        )
        if end and end < start:
            raise ValueError(
                f"Assignment {assignment_id}: effective_to before effective_from"
            )

        assignments.append({
            "assignment_id": assignment_id,
            "employee_id": employee_id,
            "effective_from": start,
            "effective_to": end,
            "department_id": department_id,
            "role_id": role_id,
            "job_level": as_int(
                row.get("job_level"), "job_level", "employee_assignments"
            ),
            "manager_id": manager_id,
            "change_reason": optional_text(row.get("change_reason")),
        })

    emails = [row["work_email"].casefold() for row in employees]
    if len(emails) != len(set(emails)):
        raise ValueError("employees: duplicate work_email values")

    employee_by_id = {row["employee_id"]: row for row in employees}
    departments_by_id = {row["department_id"] for row in departments}

    users = []
    for row in users_raw:
        user_id = required_text(row.get("user_id"), "user_id", "users")
        employee_id = required_text(
            row.get("employee_id"), "employee_id", "users"
        )

        if employee_id not in employee_by_id:
            raise ValueError(
                f"User {user_id}: unknown employee {employee_id}"
            )

        username = required_text(row.get("username"), "username", "users")
        expected_username = employee_by_id[employee_id]["work_email"]

        if username.casefold() != expected_username.casefold():
            raise ValueError(
                f"User {user_id}: username does not match employee work_email"
            )

        users.append({
            "user_id": user_id,
            "employee_id": employee_id,
            "username": username,
            "role": required_text(row.get("role"), "role", "users"),
            "is_active": as_bool(row.get("is_active"), "is_active", "users"),
            "mfa_enabled": as_bool(
                row.get("mfa_enabled"), "mfa_enabled", "users"
            ),
            "last_login_at": as_datetime(
                row.get("last_login_at"), "last_login_at", "users"
            ),
            "created_at": as_datetime(
                row.get("created_at"), "created_at", "users"
            ),
            "password_policy_note": optional_text(
                row.get("password_policy_note")
            ),
        })

    usernames = [row["username"].casefold() for row in users]
    if len(usernames) != len(set(usernames)):
        raise ValueError("users: duplicate username values")

    salary_components = []
    for row in salary_components_raw:
        salary_components.append({
            "component_id": required_text(
                row.get("component_id"), "component_id", "salary_components"
            ),
            "component_name": required_text(
                row.get("component_name"), "component_name",
                "salary_components",
            ),
            "component_type": required_text(
                row.get("component_type"), "component_type",
                "salary_components",
            ),
            "taxable": as_bool(
                row.get("taxable"), "taxable", "salary_components"
            ),
            "calculation_rule": optional_text(row.get("calculation_rule")),
            "is_fixed": as_bool(
                row.get("is_fixed"), "is_fixed", "salary_components"
            ),
        })

    holidays = []
    for row in holidays_raw:
        holidays.append({
            "holiday_date": as_date(
                row.get("holiday_date"),
                "holiday_date",
                "holidays",
                required=True,
            ),
            "holiday_name": required_text(
                row.get("holiday_name"), "holiday_name", "holidays"
            ),
            "holiday_type": optional_text(row.get("holiday_type")),
            "weekday": optional_text(row.get("weekday")),
            "falls_on_weekend": as_bool(
                row.get("falls_on_weekend"), "falls_on_weekend", "holidays"
            ),
            "calendar_note": optional_text(row.get("calendar_note")),
        })

    leave_types = []
    for row in leave_types_raw:
        quota = as_int(
            row.get("annual_quota_days"),
            "annual_quota_days",
            "leave_types",
            required=True,
        )
        carry_forward_days = as_int(
            row.get("max_carry_forward_days"),
            "max_carry_forward_days",
            "leave_types",
            required=True,
        )

        if quota < 0 or carry_forward_days < 0:
            raise ValueError(
                "leave_types: quota and carry-forward days cannot be negative"
            )

        leave_types.append({
            "leave_type_code": required_text(
                row.get("leave_type_code"), "leave_type_code", "leave_types"
            ),
            "leave_type_name": required_text(
                row.get("leave_type_name"), "leave_type_name", "leave_types"
            ),
            "pay_type": required_text(
                row.get("pay_type"), "pay_type", "leave_types"
            ),
            "annual_quota_days": quota,
            "prorated_for_joiners": as_bool(
                row.get("prorated_for_joiners"),
                "prorated_for_joiners",
                "leave_types",
            ),
            "carry_forward_allowed": as_bool(
                row.get("carry_forward_allowed"),
                "carry_forward_allowed",
                "leave_types",
            ),
            "max_carry_forward_days": carry_forward_days,
            "policy_note": optional_text(row.get("policy_note")),
        })

    return (
        departments,
        roles,
        employees,
        assignments,
        users,
        salary_components,
        holidays,
        leave_types,
    )


def upsert_rows(connection, model, rows, conflict_key, update_fields):
    if not rows:
        return

    table = model.__table__
    statement = insert(table).values(rows)
    statement = statement.on_conflict_do_update(
        index_elements=[conflict_key],
        set_={
            field: getattr(statement.excluded, field)
            for field in update_fields
        },
    )
    connection.execute(statement)


def seed_database(
    departments,
    roles,
    employees,
    assignments,
    users,
    salary_components,
    holidays,
    leave_types,
):
    engine = create_engine(settings.database_url)

    try:
        with engine.begin() as connection:
            upsert_rows(
                connection,
                Department,
                departments,
                "department_id",
                ["department_name", "cost_center", "description"],
            )

            upsert_rows(
                connection,
                Role,
                roles,
                "role_id",
                [
                    "department_id", "role_title", "level", "grade",
                    "band_min_annual", "band_mid_annual", "band_max_annual",
                    "overtime_eligible", "band_year",
                ],
            )

            # First pass: ensure employees exist before setting manager links.
            first_pass = [
                {**row, "manager_id": None}
                for row in employees
            ]
            upsert_rows(
                connection,
                Employee,
                first_pass,
                "employee_id",
                [
                    "first_name", "last_name", "full_name", "work_email",
                    "gender", "date_of_birth", "hire_date", "exit_date",
                    "employment_status", "employment_type", "department_id",
                    "role_id", "manager_id", "work_location", "work_mode",
                    "highest_education", "prior_experience_years",
                    "career_gap_months", "commute_distance_km",
                    "pf_applicable", "esi_applicable", "notice_period_days",
                    "resignation_date", "exit_type", "exit_reason",
                ],
            )

            # Restore manager relationships for workbook employee IDs only.
            employee_table = Employee.__table__
            for row in employees:
                connection.execute(
                    employee_table.update()
                    .where(
                        employee_table.c.employee_id == row["employee_id"]
                    )
                    .values(manager_id=row["manager_id"])
                )

            upsert_rows(
                connection,
                EmployeeAssignment,
                assignments,
                "assignment_id",
                [
                    "employee_id", "effective_from", "effective_to",
                    "department_id", "role_id", "job_level", "manager_id",
                    "change_reason",
                ],
            )

            # These user fields match the current User model.
            # No password_hash column exists in that model.
            upsert_rows(
                connection,
                User,
                users,
                "user_id",
                [
                    "employee_id", "username", "role", "is_active",
                    "mfa_enabled", "last_login_at", "password_policy_note",
                ],
            )

            upsert_rows(
                connection,
                SalaryComponent,
                salary_components,
                "component_id",
                [
                    "component_name", "component_type", "taxable",
                    "calculation_rule", "is_fixed",
                ],
            )

            upsert_rows(
                connection,
                Holiday,
                holidays,
                "holiday_date",
                [
                    "holiday_name", "holiday_type", "weekday",
                    "falls_on_weekend", "calendar_note",
                ],
            )

            upsert_rows(
                connection,
                LeaveType,
                leave_types,
                "leave_type_code",
                [
                    "leave_type_name", "pay_type", "annual_quota_days",
                    "prorated_for_joiners", "carry_forward_allowed",
                    "max_carry_forward_days", "policy_note",
                ],
            )

            models = {
                "departments": Department,
                "roles": Role,
                "employees": Employee,
                "employee_assignments": EmployeeAssignment,
                "users": User,
                "salary_components": SalaryComponent,
                "holidays": Holiday,
                "leave_types": LeaveType,
            }

            counts = {
                name: connection.scalar(
                    select(func.count()).select_from(model.__table__)
                )
                for name, model in models.items()
            }

        return counts
    finally:
        engine.dispose()


def main():
    parser = argparse.ArgumentParser(
        description="Validate or seed LedgerForge master workbook data."
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate workbook data without changing the database.",
    )
    args = parser.parse_args()

    datasets = load_and_validate()
    (
        departments,
        roles,
        employees,
        assignments,
        users,
        salary_components,
        holidays,
        leave_types,
    ) = datasets

    print("Workbook validation passed.")
    print(f"Departments: {len(departments)}")
    print(f"Roles: {len(roles)}")
    print(f"Employees: {len(employees)}")
    print(f"Employee assignments: {len(assignments)}")
    print(f"Users: {len(users)}")
    print(f"Salary components: {len(salary_components)}")
    print(f"Holidays: {len(holidays)}")
    print(f"Leave types: {len(leave_types)}")

    if args.dry_run:
        print("Dry run complete. Database was not changed.")
        return

    counts = seed_database(*datasets)
    print("Database seed completed successfully.")
    for table, count in counts.items():
        print(f"{table}: {count} rows currently in database")


if __name__ == "__main__":
    main()