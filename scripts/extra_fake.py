
import csv
from datetime import date, datetime
from pathlib import Path

from faker import Faker
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = BASE_DIR / "01_LedgerForge_Python_Master_Dataset.xlsx"
OUTPUT_DIR = BASE_DIR / "data"
OUTPUT_PATH = OUTPUT_DIR / "extra_fake_employees.csv"
DEMO_TODAY = date(2026, 9, 30)
EXTRA_COUNT = 50

EMPLOYEE_FIELDS = [
    "employee_id", "first_name", "last_name", "full_name", "work_email",
    "gender", "date_of_birth", "hire_date", "exit_date",
    "employment_status", "employment_type", "department_id", "role_id",
    "manager_id", "work_location", "work_mode", "highest_education",
    "prior_experience_years", "career_gap_months", "commute_distance_km",
    "pf_applicable", "esi_applicable", "notice_period_days",
    "resignation_date", "exit_type", "exit_reason", "created_at",
]


def read_records(workbook, sheet_name):
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    return [
        dict(zip(headers, row))
        for row in rows
        if any(value is not None for value in row)
    ]


def to_iso(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def main():
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    Faker.seed(2026)
    fake = Faker("en_IN")

    workbook = load_workbook(
        WORKBOOK_PATH, read_only=True, data_only=True
    )
    try:
        employees = read_records(workbook, "employees")
        departments = read_records(workbook, "departments")
        roles = read_records(workbook, "roles")
    finally:
        workbook.close()

    if not employees or not departments or not roles:
        raise ValueError("Employees, departments, and roles must not be empty.")

    department_ids = [
        str(row["department_id"])
        for row in departments
        if row.get("department_id")
    ]
    role_rows = [
        row for row in roles
        if row.get("role_id") and row.get("department_id")
    ]
    existing_ids = {
        str(row["employee_id"])
        for row in employees
        if row.get("employee_id")
    }
    existing_emails = {
        str(row["work_email"]).strip().lower()
        for row in employees
        if row.get("work_email")
    }

    if len(existing_ids) != len(employees):
        raise ValueError("Workbook contains duplicate employee IDs.")

    # Start after the largest numeric EMP- identifier.
    numeric_ids = [
        int(employee_id.split("-")[1])
        for employee_id in existing_ids
        if employee_id.startswith("EMP-")
        and employee_id.split("-")[1].isdigit()
    ]
    next_id = max(numeric_ids, default=0) + 1

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    generated = []

    for number in range(EXTRA_COUNT):
        first_name = fake.first_name()
        last_name = fake.last_name()
        full_name = f"{first_name} {last_name}"
        employee_id = f"EMP-{next_id + number:04d}"

        email_base = (
            f"{first_name}.{last_name}.{employee_id.lower()}"
            .replace(" ", "")
            .lower()
        )
        work_email = f"{email_base}@example.com"

        while work_email in existing_emails:
            work_email = (
                f"{email_base}.{fake.random_int(1000, 9999)}@example.com"
            )
        existing_emails.add(work_email)

        department_id = fake.random_element(department_ids)
        eligible_roles = [
            row for row in role_rows
            if str(row["department_id"]) == department_id
        ]
        if not eligible_roles:
            eligible_roles = role_rows

        selected_role = fake.random_element(eligible_roles)
        hire_date = fake.date_between(
            start_date=date(2018, 1, 1),
            end_date=DEMO_TODAY,
        )
        birth_date = fake.date_of_birth(
            minimum_age=21, maximum_age=60
        )
        created_at = datetime.combine(hire_date, datetime.min.time())

        generated.append({
            "employee_id": employee_id,
            "first_name": first_name,
            "last_name": last_name,
            "full_name": full_name,
            "work_email": work_email,
            "gender": None,
            "date_of_birth": birth_date.isoformat(),
            "hire_date": hire_date.isoformat(),
            "exit_date": None,
            "employment_status": "Active",
            "employment_type": "Full-time",
            "department_id": department_id,
            "role_id": str(selected_role["role_id"]),
            "manager_id": None,
            "work_location": fake.city(),
            "work_mode": fake.random_element(
                ["On-site", "Hybrid", "Remote"]
            ),
            "highest_education": fake.random_element(
                ["Bachelor's", "Master's", "Diploma"]
            ),
            "prior_experience_years": fake.random_int(0, 15),
            "career_gap_months": 0,
            "commute_distance_km": round(
                fake.pyfloat(min_value=0, max_value=40, right_digits=2),
                2,
            ),
            "pf_applicable": "Y",
            "esi_applicable": "N",
            "notice_period_days": 30,
            "resignation_date": None,
            "exit_type": None,
            "exit_reason": None,
            "created_at": created_at.isoformat(),
        })

    # Write a consistent column order matching the workbook.
    with OUTPUT_PATH.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=EMPLOYEE_FIELDS)
        writer.writeheader()
        writer.writerows(generated)

    print("Synthetic employee generation completed.")
    print(f"Seed: 2026")
    print(f"Generated records: {len(generated)}")
    print(f"Output: {OUTPUT_PATH}")
    print("Master workbook and database were not modified.")


if __name__ == "__main__":
    main()
